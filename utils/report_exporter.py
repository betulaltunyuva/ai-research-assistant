from datetime import datetime
from html import escape
from pathlib import Path
import re

from docx import Document

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import (
    getSampleStyleSheet,
    ParagraphStyle
)
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)


EXPORT_DIR = Path(
    "outputs/exports"
)

EXPORT_DIR.mkdir(
    parents=True,
    exist_ok=True
)


def safe_filename(text):
    text = text.lower()

    replacements = {
        "ç": "c",
        "ğ": "g",
        "ı": "i",
        "ö": "o",
        "ş": "s",
        "ü": "u"
    }

    for old, new in replacements.items():
        text = text.replace(
            old,
            new
        )

    text = re.sub(
        r"[^a-z0-9]+",
        "-",
        text
    )

    text = text.strip("-")

    return text[:60] or "research-report"


def clean_markdown(text):
    text = text.replace(
        "**",
        ""
    )

    text = text.replace(
        "__",
        ""
    )

    text = text.replace(
        "`",
        ""
    )

    return text.strip()


def get_pdf_fonts():
    regular_candidates = [
        Path(
            r"C:\Windows\Fonts\arial.ttf"
        ),
        Path(
            r"C:\Windows\Fonts\calibri.ttf"
        )
    ]

    bold_candidates = [
        Path(
            r"C:\Windows\Fonts\arialbd.ttf"
        ),
        Path(
            r"C:\Windows\Fonts\calibrib.ttf"
        )
    ]

    regular_path = next(
        (
            path
            for path in regular_candidates
            if path.exists()
        ),
        None
    )

    bold_path = next(
        (
            path
            for path in bold_candidates
            if path.exists()
        ),
        None
    )

    if regular_path:
        pdfmetrics.registerFont(
            TTFont(
                "ResearchFont",
                str(regular_path)
            )
        )

        if bold_path:
            pdfmetrics.registerFont(
                TTFont(
                    "ResearchFontBold",
                    str(bold_path)
                )
            )

        else:
            pdfmetrics.registerFont(
                TTFont(
                    "ResearchFontBold",
                    str(regular_path)
                )
            )

        return (
            "ResearchFont",
            "ResearchFontBold"
        )

    return (
        "Helvetica",
        "Helvetica-Bold"
    )


def export_docx(
    markdown_text,
    output_path
):
    document = Document()

    for raw_line in markdown_text.splitlines():

        line = raw_line.strip()

        if not line:
            document.add_paragraph()
            continue

        if line.startswith("### "):
            document.add_heading(
                clean_markdown(
                    line[4:]
                ),
                level=3
            )

        elif line.startswith("## "):
            document.add_heading(
                clean_markdown(
                    line[3:]
                ),
                level=2
            )

        elif line.startswith("# "):
            document.add_heading(
                clean_markdown(
                    line[2:]
                ),
                level=1
            )

        elif line.startswith("- "):
            document.add_paragraph(
                clean_markdown(
                    line[2:]
                ),
                style="List Bullet"
            )

        elif re.match(
            r"^\d+\.\s+",
            line
        ):
            text = re.sub(
                r"^\d+\.\s+",
                "",
                line
            )

            document.add_paragraph(
                clean_markdown(
                    text
                ),
                style="List Number"
            )

        else:
            document.add_paragraph(
                clean_markdown(
                    line
                )
            )

    document.save(
        str(output_path)
    )


def export_pdf(
    markdown_text,
    output_path
):
    regular_font, bold_font = (
        get_pdf_fonts()
    )

    styles = getSampleStyleSheet()

    body_style = ParagraphStyle(
        "ResearchBody",
        parent=styles["BodyText"],
        fontName=regular_font,
        fontSize=10.5,
        leading=15,
        spaceAfter=8,
        wordWrap="CJK"
    )

    title_style = ParagraphStyle(
        "ResearchTitle",
        parent=styles["Title"],
        fontName=bold_font,
        fontSize=20,
        leading=25,
        spaceAfter=18
    )

    heading2_style = ParagraphStyle(
        "ResearchHeading2",
        parent=styles["Heading2"],
        fontName=bold_font,
        fontSize=15,
        leading=19,
        spaceBefore=12,
        spaceAfter=8
    )

    heading3_style = ParagraphStyle(
        "ResearchHeading3",
        parent=styles["Heading3"],
        fontName=bold_font,
        fontSize=12,
        leading=16,
        spaceBefore=9,
        spaceAfter=6
    )

    pdf = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm
    )

    story = []

    for raw_line in markdown_text.splitlines():

        line = raw_line.strip()

        if not line:
            story.append(
                Spacer(
                    1,
                    6
                )
            )
            continue

        if line.startswith("# "):
            story.append(
                Paragraph(
                    escape(
                        clean_markdown(
                            line[2:]
                        )
                    ),
                    title_style
                )
            )

        elif line.startswith("## "):
            story.append(
                Paragraph(
                    escape(
                        clean_markdown(
                            line[3:]
                        )
                    ),
                    heading2_style
                )
            )

        elif line.startswith("### "):
            story.append(
                Paragraph(
                    escape(
                        clean_markdown(
                            line[4:]
                        )
                    ),
                    heading3_style
                )
            )

        elif line.startswith("- "):
            story.append(
                Paragraph(
                    "• "
                    + escape(
                        clean_markdown(
                            line[2:]
                        )
                    ),
                    body_style
                )
            )

        else:
            story.append(
                Paragraph(
                    escape(
                        clean_markdown(
                            line
                        )
                    ),
                    body_style
                )
            )

    pdf.build(
        story
    )


def export_report(
    query,
    final_answer
):
    timestamp = datetime.now().strftime(
        "%Y-%m-%d_%H-%M-%S"
    )

    name = safe_filename(
        query
    )

    base_name = (
        f"{timestamp}_{name}"
    )

    docx_path = (
        EXPORT_DIR /
        f"{base_name}.docx"
    )

    pdf_path = (
        EXPORT_DIR /
        f"{base_name}.pdf"
    )

    export_docx(
        final_answer,
        docx_path
    )

    export_pdf(
        final_answer,
        pdf_path
    )

    return {
        "docx_path": str(
            docx_path
        ),
        "pdf_path": str(
            pdf_path
        )
    }